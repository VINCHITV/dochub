"""
backend/tests/test_refined_prd.py

Tests for POST /projects/{project_id}/upload-refined-prd.

Coverage:
- 200 OK when status is PRD_GENERATED — prd_json replaced, status stays PRD_GENERATED.
- 200 OK when status is PRD_APPROVED — prd_json replaced, status reset to PRD_GENERATED.
- 409 when status is STORIES_GENERATED (or any post-approval state).
- 422 when an non-.docx file is uploaded.
- 404 when the project_id does not exist.
- 422 when the .docx has no recognisable Heading 1 sections.
- _build_section_dict unit tests for each section key.

Mocking strategy:
- generate_section is patched at app.routes.projects where it is imported,
  so the AI API is never called.
- get_chroma_client and get_or_create_collection are patched so ChromaDB is
  never opened (avoids HNSW file I/O in tests).
- HybridRetriever.retrieve is patched to return an empty list (no RAG nodes),
  which causes rag_context to be "(No prior knowledge base context.)" — valid
  for open_questions generation.
- All DB operations use the in-memory SQLite test engine via the test_client
  fixture from conftest.py.
"""

from __future__ import annotations

import io
import json
from typing import Any
from unittest.mock import AsyncMock, MagicMock, patch

import pytest
from docx import Document
from fastapi.testclient import TestClient
from sqlmodel import Session

from app.models import Project
from app.services.ai import GapEntry, OpenQuestionsSection
from app.services.workflow import WorkflowStatus
from app.routes.projects import _build_section_dict


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _make_docx(sections: dict[str, str]) -> bytes:
    """
    Build a minimal .docx in memory with Heading 1 paragraphs for each key.

    The section keys must be heading strings that parse_docx_by_headings can
    map via KNOWN_SECTIONS.  For tests we use the canonical aliases defined in
    docx_parser.KNOWN_SECTIONS.

    Parameters
    ----------
    sections : dict[str, str]
        Mapping of heading text -> body text.  Order is preserved.

    Returns
    -------
    bytes
        Raw .docx bytes suitable for use as an UploadFile payload.
    """
    doc = Document()
    for heading, body in sections.items():
        doc.add_heading(heading, level=1)
        if body:
            doc.add_paragraph(body)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_open_questions_result() -> OpenQuestionsSection:
    """Return a minimal OpenQuestionsSection for mocking generate_section."""
    return OpenQuestionsSection(
        type1_conflicts=[],
        type2_gaps=[
            GapEntry(question="What is the rollout timeline?", transcript_excerpt="")
        ],
        source_doc_ids=[],
    )


def _create_project(db: Session, status: WorkflowStatus) -> Project:
    """Insert a minimal Project row with the given status and return it."""
    project = Project(name="Test Project", status=status)
    db.add(project)
    db.commit()
    db.refresh(project)
    return project


# ---------------------------------------------------------------------------
# Shared patch context — applied to every test via a pytest fixture.
# Patches: ChromaDB factory calls + HybridRetriever.retrieve
# generate_section is patched per-test (its return value varies).
# ---------------------------------------------------------------------------


@pytest.fixture(autouse=True)
def _patch_rag(monkeypatch: pytest.MonkeyPatch) -> None:
    """
    Prevent any real ChromaDB or OpenAI embedding calls during the test suite.

    HybridRetriever.retrieve returns [] so rag_context falls through to the
    "(No prior knowledge base context.)" fallback string.
    """
    mock_collection = MagicMock()
    mock_client = MagicMock()

    monkeypatch.setattr(
        "app.routes.projects.get_chroma_client",
        lambda: mock_client,
    )
    monkeypatch.setattr(
        "app.routes.projects.get_or_create_collection",
        lambda client: mock_collection,
    )
    monkeypatch.setattr(
        "app.routes.projects.HybridRetriever.retrieve",
        lambda self, *args, **kwargs: [],
    )


# ---------------------------------------------------------------------------
# Test: 200 — PRD_GENERATED status, prd_json replaced, status unchanged
# ---------------------------------------------------------------------------


def test_upload_refined_prd_prd_generated_replaces_prd_json(
    test_client: TestClient,
    db_session: Session,
) -> None:
    """
    Uploading a valid .docx when status is PRD_GENERATED replaces prd_json
    and returns 200.  Status remains PRD_GENERATED (not reset, as there is
    nothing to reset from).
    """
    project = _create_project(db_session, WorkflowStatus.PRD_GENERATED)
    project_id = project.id

    docx_bytes = _make_docx({
        "Description": "We need a better onboarding flow.",
        "Problem Statement": "Users churn at the first step.",
        "Why": "Reducing churn increases ARR.",
    })

    expected_oq = _make_open_questions_result()

    with patch(
        "app.routes.projects.generate_section",
        new=AsyncMock(return_value=(expected_oq, 100, 50)),
    ):
        resp = test_client.post(
            f"/projects/{project_id}/upload-refined-prd",
            files={"file": ("refined.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )

    assert resp.status_code == 200, resp.text
    body = resp.json()

    assert body["project_id"] == project_id
    assert body["status"] == WorkflowStatus.PRD_GENERATED.value

    # open_questions must be the AI-generated result, not from DOCX
    oq = body["open_questions"]
    assert oq["type1_conflicts"] == []
    assert len(oq["type2_gaps"]) == 1
    assert oq["type2_gaps"][0]["question"] == "What is the rollout timeline?"

    # prd_json must have been written to the DB — verify via another GET
    get_resp = test_client.get(f"/projects/{project_id}")
    assert get_resp.status_code == 200
    prd_json = get_resp.json()["prd_json"]
    assert prd_json is not None
    # description section should contain the uploaded body text
    assert "onboarding" in prd_json["description"]["overview"]


# ---------------------------------------------------------------------------
# Test: 200 — PRD_APPROVED status resets to PRD_GENERATED
# ---------------------------------------------------------------------------


def test_upload_refined_prd_approved_resets_to_generated(
    test_client: TestClient,
    db_session: Session,
) -> None:
    """
    When the project is in PRD_APPROVED, a refined PRD upload resets it to
    PRD_GENERATED so the user must re-approve before generating stories.
    """
    project = _create_project(db_session, WorkflowStatus.PRD_APPROVED)
    project_id = project.id

    docx_bytes = _make_docx({"Description": "Updated description after PM review."})

    expected_oq = _make_open_questions_result()

    with patch(
        "app.routes.projects.generate_section",
        new=AsyncMock(return_value=(expected_oq, 80, 40)),
    ):
        resp = test_client.post(
            f"/projects/{project_id}/upload-refined-prd",
            files={"file": ("refined.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
        )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    # Status must be reset from PRD_APPROVED → PRD_GENERATED
    assert body["status"] == WorkflowStatus.PRD_GENERATED.value

    # Confirm the DB row was actually updated
    get_resp = test_client.get(f"/projects/{project_id}")
    assert get_resp.json()["status"] == WorkflowStatus.PRD_GENERATED.value


# ---------------------------------------------------------------------------
# Test: 409 — STORIES_GENERATED and later states are rejected
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "invalid_status",
    [
        WorkflowStatus.STORIES_GENERATED,
        WorkflowStatus.JIRA_PUSH_PENDING,
        WorkflowStatus.JIRA_PUSH_SUCCESS,
        WorkflowStatus.COMPLETED,
        WorkflowStatus.TRANSCRIPT_UPLOADED,
    ],
)
def test_upload_refined_prd_wrong_status_returns_409(
    test_client: TestClient,
    db_session: Session,
    invalid_status: WorkflowStatus,
) -> None:
    """
    Uploading a refined PRD when the project is in any status other than
    PRD_GENERATED or PRD_APPROVED must return 409 Conflict.
    """
    project = _create_project(db_session, invalid_status)
    project_id = project.id

    docx_bytes = _make_docx({"Description": "Some text."})

    resp = test_client.post(
        f"/projects/{project_id}/upload-refined-prd",
        files={"file": ("refined.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert resp.status_code == 409, resp.text
    assert invalid_status.value in resp.json()["detail"]


# ---------------------------------------------------------------------------
# Test: 422 — non-.docx file extension rejected
# ---------------------------------------------------------------------------


@pytest.mark.parametrize(
    "bad_filename",
    ["transcript.txt", "prd.pdf", "notes.md", "data.json", "file"],
)
def test_upload_refined_prd_non_docx_returns_422(
    test_client: TestClient,
    db_session: Session,
    bad_filename: str,
) -> None:
    """
    Any file that does not have a .docx extension must return 422 before
    the file content is read or the DB is queried.  This is a fast-fail guard.
    """
    project = _create_project(db_session, WorkflowStatus.PRD_GENERATED)
    project_id = project.id

    resp = test_client.post(
        f"/projects/{project_id}/upload-refined-prd",
        files={"file": (bad_filename, b"not a docx", "application/octet-stream")},
    )

    assert resp.status_code == 422, resp.text
    assert "docx" in resp.json()["detail"].lower()


# ---------------------------------------------------------------------------
# Test: 404 — project not found
# ---------------------------------------------------------------------------


def test_upload_refined_prd_project_not_found_returns_404(
    test_client: TestClient,
) -> None:
    """404 is returned when the project_id does not exist in the DB."""
    docx_bytes = _make_docx({"Description": "Some text."})

    resp = test_client.post(
        "/projects/nonexistent-project-id/upload-refined-prd",
        files={"file": ("refined.docx", docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert resp.status_code == 404, resp.text
    assert "nonexistent-project-id" in resp.json()["detail"]


# ---------------------------------------------------------------------------
# Test: 422 — .docx with no Heading 1 sections
# ---------------------------------------------------------------------------


def test_upload_refined_prd_empty_docx_returns_422(
    test_client: TestClient,
    db_session: Session,
) -> None:
    """
    A .docx that contains no recognisable Heading 1 paragraphs must return 422.
    This guards against uploading a plain paragraph document instead of a
    DocHub-structured PRD export.
    """
    project = _create_project(db_session, WorkflowStatus.PRD_GENERATED)
    project_id = project.id

    # Build a docx with only Normal-style paragraphs (no headings)
    doc = Document()
    doc.add_paragraph("This is just plain text with no headings.")
    buf = io.BytesIO()
    doc.save(buf)
    plain_docx_bytes = buf.getvalue()

    resp = test_client.post(
        f"/projects/{project_id}/upload-refined-prd",
        files={"file": ("refined.docx", plain_docx_bytes, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )

    assert resp.status_code == 422, resp.text
    assert "Heading 1" in resp.json()["detail"] or "heading" in resp.json()["detail"].lower()


# ---------------------------------------------------------------------------
# Unit tests: _build_section_dict
# ---------------------------------------------------------------------------


class TestBuildSectionDict:
    """Unit tests for the _build_section_dict helper in projects.py."""

    def test_title_no_subtitle(self) -> None:
        result = _build_section_dict("title", "Payments Redesign")
        assert result == {"title": "Payments Redesign", "subtitle": None}

    def test_title_with_subtitle(self) -> None:
        result = _build_section_dict("title", "Payments Redesign — Phase 2")
        assert result["title"] == "Payments Redesign"
        assert result["subtitle"] == "Phase 2"

    def test_description_shape(self) -> None:
        result = _build_section_dict("description", "A new checkout flow.")
        assert result["overview"] == "A new checkout flow."
        assert result["source_doc_ids"] == []

    def test_problem_with_bullets(self) -> None:
        text = "Core problem.\n- Pain point A\n- Pain point B"
        result = _build_section_dict("problem", text)
        assert result["problem_statement"] == "Core problem."
        assert "Pain point A" in result["pain_points"]
        assert "Pain point B" in result["pain_points"]
        assert result["source_doc_ids"] == []

    def test_problem_no_bullets(self) -> None:
        result = _build_section_dict("problem", "Users cannot reset passwords.")
        assert result["problem_statement"] == "Users cannot reset passwords."
        assert result["pain_points"] == []

    def test_why_two_paragraphs(self) -> None:
        text = "Rationale paragraph.\n\nBusiness value paragraph."
        result = _build_section_dict("why", text)
        assert result["rationale"] == "Rationale paragraph."
        assert result["business_value"] == "Business value paragraph."
        assert result["source_doc_ids"] == []

    def test_why_single_paragraph(self) -> None:
        result = _build_section_dict("why", "Only one paragraph.")
        assert result["rationale"] == "Only one paragraph."
        assert result["business_value"] == ""

    def test_success_bullets_become_metrics(self) -> None:
        text = "- 10% increase in conversion\n- < 200ms p99 latency"
        result = _build_section_dict("success", text)
        assert "10% increase in conversion" in result["metrics"]
        assert "< 200ms p99 latency" in result["metrics"]
        assert result["kpis"] == []
        assert result["source_doc_ids"] == []

    def test_audience_primary_only(self) -> None:
        result = _build_section_dict("audience", "Enterprise buyers.")
        assert result["primary_audience"] == "Enterprise buyers."
        assert result["secondary_audience"] is None
        assert result["personas"] == []

    def test_audience_primary_and_secondary(self) -> None:
        text = "Enterprise buyers.\n\nDeveloper advocates."
        result = _build_section_dict("audience", text)
        assert result["primary_audience"] == "Enterprise buyers."
        assert result["secondary_audience"] == "Developer advocates."

    def test_unknown_key_fallback(self) -> None:
        result = _build_section_dict("custom_section", "Some custom text.")
        assert result["content"] == "Some custom text."
        assert result["source_doc_ids"] == []
