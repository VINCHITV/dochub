"""
backend/app/services/docx_parser.py

Parses .docx files that were previously exported by docx_builder.py (or that
follow the same heading-based section layout).

Two entry points:
  parse_docx_by_headings()    — section-level extraction via Heading 1 paragraphs
  extract_embedded_metadata() — hidden [DOCHUB_METADATA] paragraph extraction

The KNOWN_SECTIONS dict maps canonical section keys to the heading strings that
the PRD template uses. The match is case-insensitive and strips leading/trailing
whitespace to tolerate minor round-trip drift.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Optional

from docx import Document

# ---------------------------------------------------------------------------
# Heading → section key mapping
# ---------------------------------------------------------------------------

KNOWN_SECTIONS: dict[str, list[str]] = {
    "title": ["title", "project title"],
    "description": ["description", "overview"],
    "problem": ["problem", "problem statement"],
    "why": ["why", "rationale", "business value"],
    "success": ["success", "success metrics", "kpis"],
    "audience": ["audience", "target audience"],
    "open_questions": ["open questions", "risks", "open questions/risks"],
}

# Prefix that marks the hidden metadata paragraph written by docx_builder.py
_METADATA_PREFIX = "[DOCHUB_METADATA]"


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def parse_docx_by_headings(content: bytes) -> dict[str, str]:
    """
    Parse a .docx document by Heading 1 paragraphs and map them to KNOWN_SECTIONS.

    Algorithm:
      1. Iterate document paragraphs in order.
      2. When a Heading 1 paragraph is found, normalise its text and look it up
         in KNOWN_SECTIONS (case-insensitive substring match).
      3. Accumulate all subsequent non-heading paragraphs as the section body
         until the next Heading 1 is encountered.

    Parameters
    ----------
    content : bytes
        Raw .docx bytes (e.g. from an uploaded file or BytesIO buffer).

    Returns
    -------
    dict[str, str]
        Mapping of section_key (e.g. "problem") -> plain text body.
        Only recognised sections are included; unknown headings are skipped.
        Returns an empty dict if the document has no Heading 1 paragraphs.

    Notes
    -----
    Non-heading paragraphs before the first Heading 1 are silently ignored.
    The [DOCHUB_METADATA] hidden paragraph is excluded from section content.
    """
    doc = Document(BytesIO(content))
    sections: dict[str, str] = {}
    current_key: Optional[str] = None
    current_lines: list[str] = []

    for para in doc.paragraphs:
        text = para.text.strip()

        if para.style.name == "Heading 1":
            # Flush previous section if any
            if current_key is not None:
                sections[current_key] = "\n".join(current_lines).strip()
            current_lines = []
            current_key = _match_section(text)
            # If heading is unrecognised, current_key is None — body lines skipped

        else:
            # Skip hidden metadata paragraph
            if text.startswith(_METADATA_PREFIX):
                continue

            if current_key is not None and text:
                current_lines.append(text)

    # Flush the last section
    if current_key is not None:
        sections[current_key] = "\n".join(current_lines).strip()

    return sections


def extract_embedded_metadata(content: bytes) -> Optional[dict]:
    """
    Find and parse the hidden [DOCHUB_METADATA] paragraph written by docx_builder.py.

    This paragraph is rendered with w:vanish so it is invisible in word processors
    but is present in the underlying XML. python-docx exposes it as a normal
    paragraph with the prefix [DOCHUB_METADATA] followed by a JSON payload.

    Parameters
    ----------
    content : bytes
        Raw .docx bytes.

    Returns
    -------
    Optional[dict]
        Parsed JSON payload (e.g. {"project_id": "...", "version": 1, ...}),
        or None if no [DOCHUB_METADATA] paragraph is found.
        Returns None (not raises) if the JSON is malformed — caller decides how
        to handle a missing/corrupt metadata paragraph.
    """
    doc = Document(BytesIO(content))

    for para in doc.paragraphs:
        text = para.text.strip()
        if text.startswith(_METADATA_PREFIX):
            json_str = text[len(_METADATA_PREFIX):]
            try:
                return json.loads(json_str)
            except json.JSONDecodeError:
                return None

    return None


# ---------------------------------------------------------------------------
# Private helpers
# ---------------------------------------------------------------------------


def _match_section(heading_text: str) -> Optional[str]:
    """
    Map a heading string to a canonical section key.

    Performs a case-insensitive substring match against each alias in
    KNOWN_SECTIONS. Returns the first matching key, or None if unrecognised.

    The substring match (rather than exact equality) handles minor heading
    variations such as "Open Questions / Risks" vs "Open Questions/Risks".
    """
    normalised = heading_text.lower().strip()
    for section_key, aliases in KNOWN_SECTIONS.items():
        for alias in aliases:
            if alias in normalised or normalised in alias:
                return section_key
    return None
